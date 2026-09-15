"""Publish client and developer audience audit revisions with configurable client identity and branding."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

from .budget import budget_markdown, budget_model
from .common import file_hash, now, read, write


def document(
    markdown: str, target: Path | str, author: str = "Drupal Upgrade Automation Toolkit"
) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    section = d.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Body Text", "List Bullet"):
        style = d.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    for name, size in [("Title", 20), ("Heading 1", 14), ("Heading 2", 11)]:
        style = d.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    for line in markdown.splitlines():
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,3}) (.*)", line)
        text = heading[2] if heading else line
        text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
        style = (
            "Title"
            if heading and len(heading[1]) == 1
            else "Heading " + str(len(heading[1]) - 1)
            if heading
            else "List Bullet"
            if line.startswith("- ")
            else None
        )
        p = d.add_paragraph(style=style)
        if style == "List Bullet":
            text = text[2:]
        for i, part in enumerate(text.split("**")):
            p.add_run(part).bold = bool(i % 2)

    for el in list(d.styles.element.iter(qn("w:pBdr"))):
        el.getparent().remove(el)
    d.core_properties.author = author
    d.save(target)


def build(
    source: Path | str,
    cfg: dict[str, Any],
    out: Path | str,
    docx: bool = False,
) -> dict[str, Any]:
    source = Path(source)
    out = Path(out)
    if out.exists() and any(out.iterdir()):
        raise ValueError("Revision output must be new; preserve historical reports")

    assessment = read(source)
    budget = budget_model(cfg)

    # Determine client branding
    client_cfg = cfg.get("client") if cfg else None
    if isinstance(client_cfg, str):
        client_name = client_cfg
        client_author = client_cfg
    elif isinstance(client_cfg, dict):
        client_name = (
            client_cfg.get("name") or client_cfg.get("id") or "Drupal Upgrade Automation Toolkit"
        )
        client_author = client_cfg.get("author") or client_name
    else:
        client_name = "Drupal Upgrade Automation Toolkit"
        client_author = "Drupal Upgrade Automation Toolkit"

    model = {
        "schemaVersion": "1.0",
        "revision": "20-hour-ai-assisted-delivery",
        "createdAt": now(),
        "releaseCandidateId": assessment["releaseCandidateId"],
        "sourceAssessment": str(source.resolve()),
        "sourceAssessmentSha256": file_hash(source),
        "releaseGate": assessment["releaseGate"],
        "findings": assessment["findings"],
        "deliveryBudget": budget,
        "client": client_name,
        "implementationAuthorized": False,
        "supersedes": "Prior estimate guidance only, including 156–296h original and 38–76h AI-assisted addendum. Historical findings and evidence remain authoritative.",
        "forecastConfidence": "low; carried forward, not recalibrated from new client scans",
        "baseline": "Earlier image-readiness failures remain diagnostic evidence. Changed capture behavior requires a new baseline; none was captured for this revision.",
    }
    write(out / "revision.json", model)
    header = f"Release readiness: **{model['releaseGate'].upper()}**. This revision changes delivery budgeting and estimate guidance. It does not authorize an upgrade or clear findings.\n\n"
    budget_text = budget_markdown(budget)
    decision = "## Delivery decision\n\nFeasibility against the toolkit default planning target remains unverified. The carried-forward 38–76h remaining-work forecast exceeds even the entire target. Use the 3h cumulative feasibility checkpoint to reconcile actual effort, blockers and scope with the engineering owner before further delivery work. Pause delivery for an explicit scope or budget decision before further work. No preparation mutation is authorized by this report.\n\n"
    assumptions = "## Forecast assumptions\n\nThe remaining-work forecast assumes AI-assisted investigation, code changes and report generation, with reviewed dependency and patch choices and complete technical QA. It excludes already completed audit work, major replacements, extensive data repair, new features, production execution and provider waiting. Earlier human effort was not reported, so total effort and remaining allowance are unknown. Business acceptance was separately provisioned at 4–8 participant hours, not measured actual effort. This is a low-confidence allowance awaiting completed scans and dependency decisions; no percentage speedup is claimed.\n\n"
    evidence = (
        "## Evidence and release requirements\n\nFive verified blockers and four evidence gaps remain, with one Green local baseline finding. Backups, code/configuration review, the full required test catalog, named exact-release business UAT, deployment rehearsal and recovery rehearsal remain mandatory. Any unfinished requirement blocks release.\n\n"
        + model["baseline"]
        + "\n\nHistorical source assessment SHA256 "
        + model["sourceAssessmentSha256"]
        + ".\n"
    )

    def rows(technical: bool = False) -> str:
        result = "## Readiness findings\n\n"
        for f in model["findings"]:
            result += (
                f"### {f['id']} {f['summary']}\n\nStatus: **{f['status']}**. "
                + f.get("technicalSummary" if technical else "clientSummary", f["summary"])
                + "\n\n"
            )
            if technical:
                result += (
                    "Action: " + f.get("recommendedAction", "See historical evidence") + "\n\n"
                )
        return result

    content = {
        "Drupal-11-Upgrade-Approval-Brief": f"# {client_name} Drupal 11 delivery approval brief\n\n"
        + header
        + budget_text
        + "\n"
        + decision
        + evidence,
        "client-summary": f"# {client_name} Drupal 11 client audit summary\n\n"
        + header
        + budget_text
        + "\n"
        + decision
        + assumptions
        + rows()
        + evidence,
        "developer-report": f"# {client_name} Drupal 11 developer audit report\n\n"
        + header
        + budget_text
        + "\n"
        + decision
        + assumptions
        + rows(True)
        + evidence,
    }
    content["combined-report"] = (
        f"# {client_name} Drupal 11 combined delivery report\n\n"
        + header
        + budget_text
        + "\n"
        + decision
        + assumptions
        + rows(True)
        + evidence
    )
    index = []
    for stem, body in content.items():
        body += "\nThis revision supersedes previous estimate guidance. Historical packets are preserved. Toolkit maintenance is outside the project delivery allowance.\n"
        (out / (stem + ".md")).write_text(body)
        files = [out / (stem + ".md")]
        if docx:
            name = "combined-upgrade-packet" if stem == "combined-report" else stem
            document(body, out / (name + ".docx"), author=client_author)
            files.append(out / (name + ".docx"))
        for p in files:
            index.append(
                {
                    "path": p.name,
                    "sha256": file_hash(p),
                    "sourceModelSha256": file_hash(out / "revision.json"),
                    "releaseGate": model["releaseGate"],
                }
            )
    write(
        out / "report-index.json",
        {"schemaVersion": "1.0", "sourceAssessmentSha256": file_hash(source), "reports": index},
    )
    return model


def validate(out: Path | str) -> dict[str, Any]:
    out = Path(out)
    index = read(out / "report-index.json")
    model = read(out / "revision.json")
    source = Path(model["sourceAssessment"])
    assert file_hash(source) == index["sourceAssessmentSha256"] == model["sourceAssessmentSha256"]
    assessment = read(source)
    assert (
        assessment["findings"] == model["findings"]
        and assessment["releaseGate"] == model["releaseGate"]
    )
    for row in index["reports"]:
        assert (
            row["sha256"] == file_hash(out / row["path"])
            and row["sourceModelSha256"] == file_hash(out / "revision.json")
            and row["releaseGate"] == model["releaseGate"]
        )
    return {
        "status": "passed",
        "reports": len(index["reports"]),
        "historicalFindingsUnchanged": True,
        "releaseGate": model["releaseGate"],
    }


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        description="Publish client/developer audience audit revisions with branding"
    )
    p.add_argument("--source")
    p.add_argument("--config")
    p.add_argument("--output", required=True)
    p.add_argument("--docx", action="store_true")
    p.add_argument("--validate", action="store_true")
    a = p.parse_args(argv)
    out = Path(a.output)
    if not a.validate:
        cfg = read(Path(a.config)) if a.config else {}
        build(Path(a.source), cfg, out, a.docx)
    print(json.dumps(validate(out)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
