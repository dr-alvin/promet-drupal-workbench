#!/usr/bin/env python3
"""Generate client, developer, combined Markdown, and optional DOCX reports.

The existing schema 3.0 packet remains authoritative. Missing audience fields are
filled deterministically from the evidence already present in the packet; they are
never interpreted as passing evidence.
"""

from __future__ import annotations

import argparse, hashlib, json, re, subprocess, sys
from datetime import datetime, timezone
from pathlib import Path

STATUSES = ("green", "amber", "red", "unknown")


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def gate_color(gate):
    return {"green": "GREEN", "amber": "AMBER", "red": "RED", "unknown": "UNKNOWN"}.get(
        gate, gate.upper()
    )


def finding_views(f):
    status = f.get("status", "unknown")
    summary = f.get("summary", "Evidence is unavailable for this check.")
    technical = f.get("technicalSummary", summary)
    client = f.get("clientSummary") or {
        "red": "This is a release blocker and must be resolved before the upgrade can proceed.",
        "amber": "This is a managed risk that needs an owner and formal approval.",
        "unknown": "The team cannot confirm this area yet because required evidence is unavailable.",
        "green": "This check has passing evidence.",
    }.get(status, "This check needs review.")
    return {
        "id": f.get("id", "F-UNNAMED"),
        "status": status,
        "summary": summary,
        "client": client,
        "technical": technical,
        "impact": f.get(
            "businessImpact",
            "May affect upgrade timing, release confidence, or recovery readiness.",
        ),
        "decision": f.get(
            "requiredDecision", "Confirm ownership and disposition before the next phase."
        ),
        "action": f.get(
            "recommendedAction", "Collect the missing evidence and record the disposition."
        ),
        "owner": f.get("ownerRole", f.get("owner", "Assigned delivery owner")),
        "estimate": f.get("estimateBand", "To be estimated after evidence review"),
        "evidence": ", ".join(f.get("evidence", [])),
    }


def phase_views(p):
    return {
        "id": p.get("id"),
        "name": p.get("name"),
        "status": p.get("status", "unknown"),
        "owner": p.get("ownerRole", p.get("owner", "Assigned delivery owner")),
        "estimate": p.get("estimateBand", "To be estimated"),
        "clientGoal": p.get(
            "clientGoal", p.get("goal", "Complete this lifecycle phase with evidence.")
        ),
        "technicalGoal": p.get(
            "technicalGoal", p.get("goal", "Complete this lifecycle phase with evidence.")
        ),
        "businessOutcome": p.get(
            "businessOutcome", "Increase confidence in a safe, supportable upgrade."
        ),
        "decisionGate": p.get(
            "decisionGate",
            "Do not advance until entry, evidence, and acceptance criteria are complete.",
        ),
        "deliverables": p.get("deliverables", p.get("evidence", [])),
        "blockedBy": p.get("blockedBy", p.get("dependencies", [])),
        "acceptance": p.get("acceptanceCriteria", []),
        "tasks": p.get("tasks", []),
        "testing": p.get("testingInstructions", []),
        "stop": p.get("stopConditions", []),
    }


def bullets(items):
    return "\n".join(f"- {x}" for x in items) if items else "- No items recorded."


def client_report(a, phases, findings):
    c = a["counts"]
    gate = gate_color(a["releaseGate"])
    lines = [
        "# Drupal 11 Upgrade Readiness — Client Summary",
        "",
        f"Release candidate: `{a['releaseCandidateId']}`",
        f"Overall decision: **{gate} — upgrade execution is not authorized**"
        if a["releaseGate"] != "green"
        else f"Overall decision: **{gate}**",
        "",
        "## What this means",
        "",
        (
            "The recorded lifecycle evidence is green for this release candidate. Execution still requires its matching approval and environment controls; this report does not authorize Production mutation."
            if a["releaseGate"] == "green"
            else "The lifecycle contains unresolved findings or missing evidence. Review the phase decisions and approved scope before authorizing further work."
        ),
        "",
        "## Status at a glance",
        "",
        f"Green: {c['green']} | Amber: {c['amber']} | Red: {c['red']} | Unknown: {c['unknown']}",
        "",
        "## Issues that need attention",
        "",
    ]
    for f in findings:
        if f["status"] in ("red", "unknown", "amber"):
            lines += [
                f"### {f['id']} — {f['status'].upper()}",
                "",
                f["client"],
                "",
                f"**Why it matters:** {f['impact']}",
                f"**Decision needed:** {f['decision']}",
                f"**Next action:** {f['action']}",
                f"**Owner:** {f['owner']}",
                "",
            ]
    lines += [
        "## Delivery phases",
        "",
        "| Phase | Status | Owner | Estimate |",
        "|---|---|---|---|",
    ]
    lines += [
        f"| {p['id']} {p['name']} | {p['status'].upper()} | {p['owner']} | {p['estimate']} |"
        for p in phases
    ]
    lines += [
        "",
        "## Approvals required",
        "",
        bullets(
            [
                "Approve the dependency and custom-code remediation plan before implementation.",
                "Confirm the non-production environment and backup/restore ownership.",
                "Approve QA coverage and mandatory business UAT for the exact release candidate.",
                "Approve deployment and restore rehearsals before any Production handoff.",
            ]
        ),
        "",
        "## Recommended next step",
        "",
        (
            "Follow the reviewed handoff and sign-off process for this exact release. No Production change is authorized by this report."
            if a["releaseGate"] == "green"
            else "Resolve recorded blockers and collect missing evidence before progressing. No Production change is authorized by this report."
        ),
    ]
    return "\n".join(lines) + "\n"


def developer_report(a, phases, findings, root):
    c = a["counts"]
    lines = [
        "# Drupal 11 Upgrade Readiness — Developer Report",
        "",
        f"Release candidate: `{a['releaseCandidateId']}`",
        f"Gate: **{gate_color(a['releaseGate'])}**",
        "",
        "## Technical disposition",
        "",
        f"Finding counts: green={c['green']}, amber={c['amber']}, red={c['red']}, unknown={c['unknown']}.",
        "",
        "## Findings",
        "",
        "| ID | Status | Technical summary | Owner | Evidence |",
        "|---|---|---|---|---|",
    ]
    lines += [
        "| {} | {} | {} | {} | {} |".format(
            f["id"], f["status"], f["technical"].replace("|", "/"), f["owner"], f["evidence"]
        )
        for f in findings
    ]
    for p in phases:
        lines += [
            "",
            f"## {p['id']}. {p['name']}",
            "",
            f"**Status:** {p['status']}  ",
            f"**Owner:** {p['owner']}  ",
            f"**Estimate:** {p['estimate']}",
            "",
            f"**Technical goal:** {p['technicalGoal']}",
            "",
            "### Tasks",
            "",
            bullets(p["tasks"]),
            "",
            "### Dependencies and blockers",
            "",
            bullets(p["blockedBy"]),
            "",
            "### Evidence/deliverables",
            "",
            bullets(p["deliverables"]),
            "",
            "### Acceptance criteria",
            "",
            bullets(p["acceptance"]),
            "",
            "### Testing instructions",
            "",
            bullets(p["testing"]),
            "",
            "### Stop conditions",
            "",
            bullets(p["stop"]),
        ]
    discovery = root / "01-upgrade-control/context.json"
    composer = root / "01-upgrade-control/composer-resolution.json"
    modules = root / "01-upgrade-control/module-readiness.csv"
    sequence = root / "01-upgrade-control/deployment-sequence.json"
    if discovery.is_file() or composer.is_file() or modules.is_file() or sequence.is_file():
        lines += ["", "## Discovery and release controls", ""]
        for path in (discovery, composer, modules, sequence):
            if path.is_file():
                lines.append(f"- `{path.relative_to(root).as_posix()}`")
        if sequence.is_file():
            seq = load(sequence)
            safety = seq.get("moduleRemovalSafety", {})
            lines += [
                "",
                f"Module-removal configuration safety: **{'pass' if safety.get('configurationBeforePackageRemoval') else 'review required'}**",
            ]
    lines += [
        "",
        "## Evidence index",
        "",
        bullets(sorted(x.relative_to(root).as_posix() for x in root.rglob("*") if x.is_file())),
    ]
    return "\n".join(lines) + "\n"


def combined_report(a, phases, findings, root):
    return (
        client_report(a, phases, findings)
        + "\n---\n\n"
        + developer_report(a, phases, findings, root)
    )


def make_docx(markdown_path, output_path, title):
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}")
    status_palette = {
        "green": "067647",
        "amber": "8A4B00",
        "red": "B42318",
        "unknown": "475467",
    }

    def clean(value):
        value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
        value = re.sub(r"`([^`]*)`", r"\1", value)
        return value

    def shade(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        node = tc_pr.find(qn("w:shd"))
        if node is None:
            node = OxmlElement("w:shd")
            tc_pr.append(node)
        node.set(qn("w:fill"), fill)

    def margins(cell, top=100, start=120, bottom=100, end=120):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tc_mar.find(qn(f"w:{side}"))
            if node is None:
                node = OxmlElement(f"w:{side}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def repeat_header(row):
        tr_pr = row._tr.get_or_add_trPr()
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "true")
        tr_pr.append(marker)

    def set_cell_text(cell, value, *, bold=False, white=False, center=False, size=9):
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(value)
        run.bold = bold
        run.font.size = Pt(size)
        if white:
            run.font.color.rgb = RGBColor(255, 255, 255)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margins(cell)

    def status_value(value):
        normalized = clean(value).strip().lower()
        return normalized if normalized in status_palette else None

    def style_status_cell(cell, value):
        status = status_value(value)
        if not status:
            return False
        shade(cell, status_palette[status])
        set_cell_text(cell, status.upper(), bold=True, white=True, center=True, size=9.5)
        return True

    def add_status_strip(counts):
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for index, status in enumerate(("green", "amber", "red", "unknown")):
            cell = table.rows[0].cells[index]
            cell.width = Inches(1.625)
            shade(cell, status_palette[status])
            set_cell_text(
                cell,
                f"{status.upper()}  {counts.get(status, 0)}",
                bold=True,
                white=True,
                center=True,
                size=9.5,
            )
        return table

    def add_status_banner(label, status, suffix=""):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        table.rows[0].cells[0].width = Inches(1.25)
        table.rows[0].cells[1].width = Inches(5.25)
        shade(table.rows[0].cells[0], status_palette[status])
        set_cell_text(
            table.rows[0].cells[0], status.upper(), bold=True, white=True, center=True, size=10
        )
        set_cell_text(table.rows[0].cells[1], f"{label}{suffix}".strip(), bold=True, size=10)
        return table

    def column_widths(headers):
        normalized = [clean(x).strip().lower() for x in headers]
        if normalized == ["id", "status", "technical summary", "owner", "evidence"]:
            return [0.55, 0.8, 2.45, 1.25, 1.45]
        if normalized == ["phase", "status", "owner", "estimate"]:
            return [2.45, 0.85, 1.65, 1.55]
        return [6.5 / len(headers)] * len(headers)

    def add_markdown_table(rows):
        widths = column_widths(rows[0])
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for row_index, values in enumerate(rows):
            row = table.add_row()
            if row_index == 0:
                repeat_header(row)
            for index, value in enumerate(values):
                cell = row.cells[index]
                cell.width = Inches(widths[index])
                if row_index == 0:
                    shade(cell, "1F4D78")
                    set_cell_text(
                        cell, clean(value), bold=True, white=True, center=index in (0, 1), size=9
                    )
                elif not style_status_cell(cell, value):
                    set_cell_text(cell, clean(value), center=index in (0, 1), size=8.5)
        return table

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    doc.core_properties.title = title
    doc.core_properties.author = "Drupal 11 Upgrade Readiness skill"
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            doc.add_paragraph()
            index += 1
            continue
        if line.startswith("# "):
            p = doc.add_heading(clean(line[2:]), 0)
        elif line.startswith("## "):
            p = doc.add_heading(clean(line[3:]), 1)
        elif line.startswith("### "):
            p = doc.add_heading(clean(line[4:]), 2)
        elif line.startswith("- "):
            p = doc.add_paragraph(clean(line[2:]), style="List Bullet")
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                candidate = lines[index].strip()
                if set(candidate.replace("|", "").replace("-", "").replace(":", "").strip()):
                    rows.append([x.strip() for x in candidate.strip("|").split("|")])
                index += 1
            if rows:
                add_markdown_table(rows)
            continue
        else:
            count_match = re.search(
                r"(?:Finding counts:\s*)?green[=:]\s*(\d+).*amber[=:]\s*(\d+).*red[=:]\s*(\d+).*unknown[=:]\s*(\d+)",
                clean(line),
                re.I,
            )
            gate_match = re.match(
                r"(?:Overall decision|Gate):\s*(green|amber|red|unknown)(.*)", clean(line), re.I
            )
            phase_status = re.match(r"Status:\s*(green|amber|red|unknown)\s*$", clean(line), re.I)
            if count_match:
                add_status_strip(
                    dict(zip(("green", "amber", "red", "unknown"), map(int, count_match.groups())))
                )
            elif gate_match:
                add_status_banner("Release gate", gate_match.group(1).lower(), gate_match.group(2))
            elif phase_status:
                add_status_banner("Phase status", phase_status.group(1).lower())
            else:
                p = doc.add_paragraph(clean(line))
        index += 1
    doc.save(output_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("artifact_root", type=Path)
    ap.add_argument("--docx", action="store_true")
    args = ap.parse_args()
    root = args.artifact_root
    a = load(root / "assessment.json")
    rawph = load(root / "phase-status.json")["phases"]
    phases = [phase_views(x) for x in rawph]
    findings = [
        finding_views(x) for x in sorted(a.get("findings", []), key=lambda x: x.get("id", ""))
    ]
    reports = {
        "client-summary.md": client_report(a, phases, findings),
        "developer-report.md": developer_report(a, phases, findings, root),
        "combined-report.md": combined_report(a, phases, findings, root),
    }
    for name, text in reports.items():
        (root / name).write_text(text, encoding="utf-8")
    if args.docx:
        make_docx(
            root / "client-summary.md",
            root / "client-summary.docx",
            "Drupal 11 Upgrade Readiness — Client Summary",
        )
        make_docx(
            root / "developer-report.md",
            root / "developer-report.docx",
            "Drupal 11 Upgrade Readiness — Developer Report",
        )
        make_docx(
            root / "combined-report.md",
            root / "combined-upgrade-packet.docx",
            "Drupal 11 Upgrade Readiness — Combined Packet",
        )
    created = datetime.now(timezone.utc).isoformat()
    idx = []
    for path in (
        sorted(root.glob("client-summary.*"))
        + sorted(root.glob("developer-report.*"))
        + sorted(root.glob("combined-*"))
    ):
        if path.name == "report-index.json":
            continue
        audience = (
            "client"
            if path.name.startswith("client-")
            else "developer"
            if path.name.startswith("developer-")
            else "combined"
        )
        idx.append(
            {
                "path": path.relative_to(root).as_posix(),
                "audience": audience,
                "releaseCandidateId": a["releaseCandidateId"],
                "sourceAssessmentSha256": sha(root / "assessment.json"),
                "generatedAt": created,
                "sha256": sha(path),
                "status": "generated",
                "gate": a["releaseGate"],
            }
        )
    (root / "report-index.json").write_text(
        json.dumps({"schemaVersion": "3.0", "reports": idx}, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"artifactRoot": str(root), "reports": len(idx), "docx": args.docx}, indent=2))


if __name__ == "__main__":
    main()
